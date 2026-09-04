# Extract table headers and code content from the docx file
import sys
import os

try:
    from docx import Document
except ImportError:
    print("ERROR: python-docx not installed")
    sys.exit(1)

docx_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 
                          "copyright", "BioChainAI_Source_Code_Report (1).docx")

print(f"Reading: {docx_path}")
doc = Document(docx_path)

print(f"\nTotal tables found: {len(doc.tables)}")
print("=" * 80)

for i, table in enumerate(doc.tables):
    header = table.rows[0].cells[0].text.strip()
    if len(table.rows) > 1:
        code = table.rows[1].cells[0].text.strip()
        code_lines = len(code.split('\n'))
    else:
        code = ""
        code_lines = 0
    print(f"Table {i+1}: {header}  ({code_lines} lines of code)")

print("=" * 80)
print(f"Total files in report: {len(doc.tables)}")
